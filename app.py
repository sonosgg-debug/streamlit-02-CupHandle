import socket
socket.setdefaulttimeout(15.0)

import streamlit as st
import pandas as pd
import numpy as np
import yfinance as yf
import FinanceDataReader as fdr
import concurrent.futures
import datetime

# 한국 표준시(KST) 타임존 (UTC+9)
KST = datetime.timezone(datetime.timedelta(hours=9))
import io
import plotly.graph_objects as go
from plotly.subplots import make_subplots

from openpyxl.utils import get_column_letter
from tickers import get_krx_tickers
from screener import screen_single_stock

STANDARD_CHART_THEME = {
    'paper_bgcolor': '#1E293B',    # Tailwind Slate-800 (외곽 카드 배경)
    'plot_bgcolor': '#0F172A',     # Tailwind Slate-900 (내부 딥 블랙 플롯)
    'text_main': '#F8FAFC',        # 타이틀/헤더 텍스트 (순백색)
    'text_body': '#E2E8F0',        # 본문 및 축 라벨 (부드러운 화이트)
    'text_muted': '#CBD5E1',       # 축 눈금 수치 텍스트 (Slate-300)
    'grid_color': '#334155',       # 그리드 격자선 (Slate-700)
    'border_color': '#475569',     # 축 기준선 (Slate-600)
    'legend_bg': 'rgba(30, 41, 59, 0.85)',
    'legend_border': '#334155',
    'hover_bg': 'rgba(15, 23, 42, 0.9)',
    'hover_border': '#334155'
}


def fmt_curr(val, ticker):
    if pd.isna(val):
        return ""
    if ticker.endswith('.KS') or ticker.endswith('.KQ'):
        return f"{val:,.0f}원"
    else:
        return f"${val:,.2f}"

# 페이지 설정
st.set_page_config(
    page_title="Cup with Handle Stock Screener",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# 커스텀 CSS로 UI 스타일링 (다크 테마 최적화 및 시인성 개선)
st.markdown("""
<style>
    /* Streamlit 고정 상단 헤더 배경 투명화 */
    header[data-testid="stHeader"] {
        background: transparent !important;
    }

    .main .block-container,
    [data-testid="stMainBlockContainer"],
    .block-container {
        padding-top: 2.0rem !important;
    }
    .main-title {
        font-size: 2.0rem !important;
        font-weight: 800 !important;
        color: #8AB4F8 !important;
        -webkit-text-fill-color: #8AB4F8 !important;
        text-align: center !important;
        margin-bottom: 0.2rem;
    }
    .sub-title {
        font-size: 0.9rem;
        color: #BDC1C6; /* 밝은 회색으로 가독성 향상 */
        text-align: center;
        margin-bottom: 2rem;
    }
    .metric-card {
        background-color: #202124; /* 검정색 계열의 배경 적용 */
        color: #F1F3F4; /* 폰트를 밝은 색상으로 강제 지정 */
        padding: 15px;
        border-radius: 8px;
        border-left: 5px solid #8AB4F8; /* 하늘색 테두리 포인트 */
        margin-bottom: 10px;
    }
    .metric-card ul, .metric-card li {
        font-size: 0.9rem;
        line-height: 1.5;
    }
    /* 안내문(st.info) 폰트 크기 및 이모지 아이콘 크기 축소 */
    .stAlert p, .stAlert [data-testid="stMarkdownContainer"] {
        font-size: 0.88rem !important;
    }
    .stAlert [data-testid="stAlertDynamicIcon"],
    .stAlert [data-testid="stIconEmoji"] {
        font-size: 1.0rem !important;
        width: 1.0rem !important;
        height: 1.0rem !important;
    }

    /* 사이드바 스타일링 */
    section[data-testid="stSidebar"], [data-testid="stSidebar"] {
        background-color: #1e293b !important;
        border-right: 1px solid #334155 !important;
    }
    section[data-testid="stSidebar"] h1,
    section[data-testid="stSidebar"] h2,
    section[data-testid="stSidebar"] h3 {
        color: #f8fafc !important;
        -webkit-text-fill-color: #f8fafc !important;
    }

    /* =========================================================
       사이드바 접기(<<) 및 펼치기(>>) 버튼 항상 표시 및 시인성/대비 강화
       ========================================================= */
    /* 1. 사이드바가 열려 있을 때 접기 버튼 (<<) 상시 표시 */
    [data-testid="stSidebarCollapseButton"] {
        visibility: visible !important;
        opacity: 1 !important;
        display: inline-flex !important;
    }
    
    [data-testid="stSidebarCollapseButton"] button {
        visibility: visible !important;
        opacity: 1 !important;
        background-color: #1e293b !important;       /* 진한 네이비 배경 */
        border: 1.5px solid #38bdf8 !important;     /* 선명한 스카이블루 테두리로 상자 명확화 */
        border-radius: 8px !important;
        width: 38px !important;
        height: 38px !important;
        display: flex !important;
        align-items: center !important;
        justify-content: center !important;
        box-shadow: 0 2px 8px rgba(0, 0, 0, 0.4), 0 0 6px rgba(56, 189, 248, 0.2) !important;
        transition: all 0.2s ease !important;
    }
    
    /* 상자 내부의 << 아이콘(Material Icon span/svg/문자)을 순백색으로 강제하여 상자와 극명한 대비 구현 */
    [data-testid="stSidebarCollapseButton"] button *,
    [data-testid="stSidebarCollapseButton"] span,
    [data-testid="stSidebarCollapseButton"] [data-testid="stIconMaterial"],
    [data-testid="stSidebarCollapseButton"] svg {
        color: #ffffff !important;
        fill: #ffffff !important;
        opacity: 1 !important;
        visibility: visible !important;
        font-size: 1.35rem !important;
        font-weight: 700 !important;
    }
    
    /* 호버(PC) 및 터치 시 반전 효과 */
    [data-testid="stSidebarCollapseButton"] button:hover {
        background-color: #38bdf8 !important;
        border-color: #38bdf8 !important;
    }
    [data-testid="stSidebarCollapseButton"] button:hover * {
        color: #0f172a !important;
        fill: #0f172a !important;
    }

    /* 2. 사이드바 헤더 영역 패딩 및 정렬 보정 */
    [data-testid="stSidebarHeader"] {
        padding-top: 0.5rem !important;
        padding-bottom: 0.5rem !important;
    }

    /* 3. 사이드바가 닫혔을 때 다시 여는 버튼 (>>) 시인성 강화 */
    [data-testid="stSidebarCollapsedControl"] {
        visibility: visible !important;
        opacity: 1 !important;
    }
    
    [data-testid="stSidebarCollapsedControl"] button {
        background-color: #1e293b !important;
        border: 1.5px solid #38bdf8 !important;
        border-radius: 8px !important;
        box-shadow: 0 2px 8px rgba(0, 0, 0, 0.4), 0 0 6px rgba(56, 189, 248, 0.2) !important;
    }
    
    [data-testid="stSidebarCollapsedControl"] button *,
    [data-testid="stSidebarCollapsedControl"] span,
    [data-testid="stSidebarCollapsedControl"] [data-testid="stIconMaterial"],
    [data-testid="stSidebarCollapsedControl"] svg {
        color: #38bdf8 !important;
        fill: #38bdf8 !important;
        opacity: 1 !important;
        visibility: visible !important;
        font-size: 1.35rem !important;
    }

    /* Primary Button Styling (39 DividendStock 테마 통일) */
    .stButton button[kind="primary"],
    .stButton > button[kind="primary"],
    section[data-testid="stSidebar"] button[kind="primary"] {
        background-color: #2563eb !important;
        color: #ffffff !important;
        border: none !important;
        font-weight: 600 !important;
        border-radius: 6px !important;
        transition: all 0.2s ease !important;
    }
    .stButton button[kind="primary"]:hover,
    .stButton > button[kind="primary"]:hover,
    section[data-testid="stSidebar"] button[kind="primary"]:hover {
        background-color: #1d4ed8 !important;
        box-shadow: 0 0 10px rgba(37, 99, 235, 0.4) !important;
    }

    /* 다운로드 버튼 공통 통일 스타일 */
    div[data-testid="stDownloadButton"] > button,
    .stDownloadButton > button {
        background-color: #334155 !important;
        color: #f8fafc !important;
        border: 1px solid #475569 !important;
        border-radius: 6px !important;
        font-size: 0.875rem !important;
        font-weight: 500 !important;
        height: 38px !important;
        min-height: 38px !important;
        max-height: 38px !important;
        line-height: 36px !important;
        padding: 0 16px !important;
        display: inline-flex !important;
        align-items: center !important;
        justify-content: center !important;
        text-align: center !important;
        transition: all 0.2s ease-in-out !important;
        box-sizing: border-box !important;
    }
    div[data-testid="stDownloadButton"] > button:hover,
    .stDownloadButton > button:hover {
        background-color: #475569 !important;
        border-color: #38bdf8 !important;
        color: #ffffff !important;
        box-shadow: 0 0 10px rgba(56, 189, 248, 0.25) !important;
    }
    div[data-testid="stDownloadButton"] > button:active,
    .stDownloadButton > button:active {
        background-color: #1e293b !important;
        border-color: #0284c7 !important;
    }
    div[data-testid="stDownloadButton"] > button p,
    div[data-testid="stDownloadButton"] > button span,
    .stDownloadButton > button p,
    .stDownloadButton > button span {
        font-size: 0.875rem !important;
        font-weight: 500 !important;
        color: inherit !important;
        line-height: inherit !important;
        margin: 0 !important;
        padding: 0 !important;
    }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-title">William O\'Neal "Cup with Handle" 패턴 스크리너</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-title">한국 및 미국 주식시장 종목 중 컵앤핸들(Cup & Handle) 돌파 및 거래량 동반 종목 발굴 프로그램</div>', unsafe_allow_html=True)

# 기법 소개
with st.expander("ℹ️ 윌리엄 오닐의 'Cup with Handle' (컵앤핸들) 패턴이란?"):
    st.markdown("""
    **컵앤핸들(Cup with Handle)** 패턴은 전설적인 성장주 투자자 **윌리엄 오닐(William O'Neal)**이 정립한 대표적인 강세 지속 패턴입니다.
    조정 국면을 거친 후 전고점을 강한 거래량과 함께 돌파할 때 매수 타이밍을 잡는 기법입니다.
    
    ### 📌 핵심 스크리닝 요건
    1. **이전 상승 추세 (Prior Trend)**:
       * 패턴 형성 시작(Peak A) 전에 최소 **25% 이상**의 뚜렷한 주가 상승세가 선행되어야 합니다.
    2. **컵 (Cup) 구간**:
       * 주가가 왼쪽 고점(Peak A)에서 하락 조정된 뒤, 둥근 U자형 바닥(Valley B)을 형성하며 점진적으로 상승하여 오른쪽 고점(Peak C)에 도달합니다.
       * **기간**: 컵 형성 기간은 최소 35거래일(약 7주)에서 최대 220거래일(약 10개월)입니다.
       * **깊이**: 최고점 대비 10% ~ 50%의 조정 비율이 일반적입니다.
    3. **핸들 (Handle) 구간**:
       * 컵 오른쪽 고점(Peak C)을 형성한 후 가볍게 되밀리는 조정 구간을 거칩니다.
       * **기간**: 최소 5거래일에서 최대 30거래일 정도 유지됩니다.
       * **깊이**: 컵 오른쪽 고점 대비 최대 15% 이내의 얕은 조정이어야 하며, 컵의 상반부(상단 50% 영역) 안에서 마무리되어야 합니다.
       * **거래량**: 핸들이 조정받는 동안 거래량이 점진적으로 줄어들며 매물이 소진되는 모습을 보여야 합니다.
    4. **돌파 (Breakout)**:
       * 핸들의 조정기를 마치고 컵의 저항선(Peak C 가격 또는 핸들 최고점)을 **종가 기준으로 상향 돌파**해야 합니다.
       * 이때 돌파 당일의 거래량은 **최근 20일 평균 거래량의 1.5배(150%) 이상** 급증하여 매수세 유입을 증명해야 합니다.
    """)

# 세션 상태 초기화
if 'screened_df' not in st.session_state:
    st.session_state.screened_df = None
if 'last_run_time' not in st.session_state:
    st.session_state.last_run_time = None
if 'market_type_used' not in st.session_state:
    st.session_state.market_type_used = None
if 'raw_screened_df' not in st.session_state:
    st.session_state.raw_screened_df = None

# 사이드바 설정 영역
with st.sidebar:
    st.markdown(
        """
        <div style='padding: 2px 0 12px 0;'>
            <div style='font-size: 1.25rem; font-weight: 700; color: #f8fafc; letter-spacing: -0.01em; display: flex; align-items: center; gap: 8px;'>
                <span>⚙️</span> 스크리닝 조건 설정
            </div>
            <div style='font-size: 0.82rem; color: #94a3b8; margin-top: 4px; line-height: 1.4;'>
                윌리엄 오닐 컵앤핸들 패턴 발굴 조건을 설정합니다.
            </div>
        </div>
        <hr style='border: 0; height: 1px; background-color: #334155; margin: 10px 0 16px 0;'>
        """,
        unsafe_allow_html=True
    )

    market_choice = st.selectbox(
        "🏛️ 시장 선택",
        ["KOSPI", "KOSDAQ", "S&P 500", "NASDAQ"],
        index=0
    )

    st.markdown("<hr style='border: 0; height: 1px; background-color: #334155; margin: 16px 0;'>", unsafe_allow_html=True)
    st.markdown("<div style='font-size: 0.95rem; font-weight: 700; color: #e2e8f0; margin-bottom: 6px;'>🎯 컵(Cup) 패턴 설정</div>", unsafe_allow_html=True)
    min_cup_width = st.slider("최소 컵 기간 (영업일)", 20, 90, 35, step=5)
    max_cup_width = st.slider("최대 컵 기간 (영업일)", 100, 300, 220, step=10)
    min_cup_depth = st.slider("최소 컵 깊이 (%)", 5, 30, 10, step=1) / 100.0
    max_cup_depth = st.slider("최대 컵 깊이 (%)", 30, 70, 50, step=5) / 100.0

    st.markdown("---")
    st.markdown("<div style='font-size: 0.95rem; font-weight: 700; color: #e2e8f0; margin-bottom: 6px;'>🎯 핸들(Handle) 패턴 설정</div>", unsafe_allow_html=True)
    min_handle_width = st.slider("최소 핸들 기간 (영업일)", 2, 15, 5, step=1)
    max_handle_width = st.slider("최대 핸들 기간 (영업일)", 15, 50, 30, step=5)
    max_handle_depth = st.slider("최대 핸들 깊이 (%)", 5, 35, 15, step=1) / 100.0

    st.markdown("---")
    st.markdown("<div style='font-size: 0.95rem; font-weight: 700; color: #e2e8f0; margin-bottom: 6px;'>🎯 돌파 및 거래량 필터</div>", unsafe_allow_html=True)
    breakout_vol_factor = st.slider("최소 돌파 거래량 배수", 1.0, 3.0, 1.5, step=0.1)
    prior_trend_gain = st.slider("최소 선행 상승률 (%)", 10, 50, 25, step=5) / 100.0
    breakout_window = st.slider("최근 돌파 허용 기간 (영업일)", 1, 15, 5, step=1)

    chunk_size = st.number_input(
        "데이터 일괄 요청 크기 (Chunk)",
        min_value=10,
        max_value=100,
        value=50,
        step=10,
        help="yfinance API로 한 번에 다운로드할 종목 개수입니다. 너무 크게 설정하면 API 에러가 발생할 수 있습니다."
    )

    # 스크리닝 시작 버튼
    start_screening = st.button("🔍 스크리닝 시작", type="primary", use_container_width=True)

if start_screening:
    market_map = {
        "KOSPI": "KOSPI",
        "KOSDAQ": "KOSDAQ",
        "S&P 500": "S&P 500",
        "NASDAQ": "NASDAQ 100",
        # 하위 호환 매핑
        "코스피 (KOSPI)": "KOSPI",
        "코스닥 (KOSDAQ)": "KOSDAQ",
        "전체 시장 (KOSPI + KOSDAQ)": "ALL",
        "미국 S&P 500 (US)": "S&P 500",
        "미국 NASDAQ 100 (US)": "NASDAQ 100"
    }
    selected_market = market_map.get(market_choice, "KOSPI")
    
    with st.spinner("상장 종목 목록을 가져오는 중..."):
        try:
            tickers_df = get_krx_tickers(selected_market)
            total_count = len(tickers_df)
            st.info(f"수집 대상 종목: 총 {total_count}개 (우선주/스팩 필터링 완료)")
        except Exception as e:
            st.error(f"종목 목록 수집 실패: {e}")
            tickers_df = pd.DataFrame()
            
    if not tickers_df.empty:
        progress_bar = st.progress(0)
        status_text = st.empty()
        
        results = []
        tickers = tickers_df['ticker'].tolist()
        name_map = dict(zip(tickers_df['ticker'], tickers_df['회사명']))
        
        total_tickers = len(tickers)
        chunks = [tickers[i:i + chunk_size] for i in range(0, total_tickers, chunk_size)]
        
        start_time = datetime.datetime.now()
        
        params = {
            'min_cup_width': min_cup_width,
            'max_cup_width': max_cup_width,
            'min_cup_depth': min_cup_depth,
            'max_cup_depth': max_cup_depth,
            'min_handle_width': min_handle_width,
            'max_handle_width': max_handle_width,
            'max_handle_depth': max_handle_depth,
            'breakout_vol_factor': breakout_vol_factor,
            'prior_trend_gain': prior_trend_gain,
            'breakout_window': breakout_window
        }
        
        start_date = (datetime.datetime.now() - datetime.timedelta(days=1095)).strftime('%Y-%m-%d')
        
        for idx, chunk in enumerate(chunks):
            status_text.text(f"데이터 다운로드 및 컵앤핸들 분석 중... [{idx+1}/{len(chunks)}] (진행률: {int((idx+1)/len(chunks)*100)}%)")
            progress_bar.progress((idx + 1) / len(chunks))
            
            kr_chunk = [t for t in chunk if t.endswith('.KS') or t.endswith('.KQ')]
            us_chunk = [t for t in chunk if not (t.endswith('.KS') or t.endswith('.KQ'))]
            
            # 1. 한국 주식: FinanceDataReader 멀티스레드 병렬 수집 (네이버 금융 공식 시세)
            if kr_chunk:
                def fetch_kr_stock(t):
                    code = t.split('.')[0]
                    try:
                        df = fdr.DataReader(code, start_date)
                        if df is not None and not df.empty:
                            df = df.dropna(subset=['Close', 'High', 'Low', 'Volume'])
                            if df.index.tz is not None:
                                df.index = df.index.tz_localize(None)
                            return t, df
                    except Exception:
                        pass
                    return t, None

                with concurrent.futures.ThreadPoolExecutor(max_workers=min(len(kr_chunk), 20)) as executor:
                    fetched_kr = list(executor.map(fetch_kr_stock, kr_chunk))

                for ticker, df_single in fetched_kr:
                    if df_single is None or len(df_single) < 250:
                        continue
                    try:
                        res = screen_single_stock(ticker, name_map[ticker], df_single, params)
                        if res:
                            results.append(res)
                    except Exception:
                        continue

            # 2. 미국 주식: yfinance 일괄 다운로드
            if us_chunk:
                try:
                    data = yf.download(us_chunk, period="3y", group_by="ticker", progress=False)
                    for ticker in us_chunk:
                        try:
                            if isinstance(data.columns, pd.MultiIndex):
                                ticker_level = 'Ticker' if 'Ticker' in data.columns.names else 1
                                tickers_in_data = data.columns.get_level_values(ticker_level).unique()
                                if ticker not in tickers_in_data:
                                    continue
                                df_single = data.xs(ticker, level=ticker_level, axis=1).dropna(subset=['Close', 'High', 'Low', 'Volume'])
                            else:
                                df_single = data.dropna(subset=['Close', 'High', 'Low', 'Volume'])
                                
                            if df_single.index.tz is not None:
                                df_single.index = df_single.index.tz_localize(None)

                            if len(df_single) < 250:
                                continue
                                
                            res = screen_single_stock(ticker, name_map[ticker], df_single, params)
                            if res:
                                results.append(res)
                        except Exception:
                            continue
                except Exception:
                    pass
                
        # 프로그레스바 초기화
        progress_bar.empty()
        status_text.empty()
        
        if results:
            df_final = pd.DataFrame(results)
            # 출력용 한글 칼럼명 매핑
            df_final_display = df_final.rename(columns={
                'symbol': '티커',
                'name': '종목명',
                'breakout_date': '돌파 감지일',
                'breakout_price': '돌파 가격',
                'cup_depth_pct': '컵 깊이(%)',
                'handle_depth_pct': '핸들 깊이(%)',
                'cup_width_days': '컵 기간(일)',
                'handle_width_days': '핸들 기간(일)',
                'volume_increase_ratio': '거래량 비율'
            })
            
            # 불필요한 내부 칼럼은 제외하고 테이블용 구성
            display_cols = ['티커', '종목명', '돌파 감지일', '돌파 가격', '컵 깊이(%)', '핸들 깊이(%)', '컵 기간(일)', '핸들 기간(일)', '거래량 비율']
            df_final_display = df_final_display[display_cols].copy()
            
            st.session_state.screened_df = df_final_display
            st.session_state.raw_screened_df = df_final  # 원본 데이터도 세션 저장
        else:
            st.session_state.screened_df = pd.DataFrame()
            st.session_state.raw_screened_df = pd.DataFrame()
            
        st.session_state.last_run_time = datetime.datetime.now(KST).strftime('%Y-%m-%d %H:%M:%S')
        st.session_state.market_type_used = market_choice

# 결과 디스플레이
if st.session_state.screened_df is not None:
    st.success(f"🔍 스크리닝 완료! (실행 시각: {st.session_state.last_run_time} | 대상: {st.session_state.market_type_used})")
    
    if st.session_state.screened_df.empty:
        st.warning("조건에 부합하는 종목이 발견되지 않았습니다. 파라미터를 조절하여 다시 스크리닝해 보세요.")
    else:
        # --- 엑셀 저장 및 다운로드 파일 사전 생성 ---
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df_excel = st.session_state.raw_screened_df.copy()
            df_excel = df_excel.rename(columns={
                'symbol': '티커',
                'name': '종목명',
                'market': '시장',
                'breakout_date': '돌파 감지일',
                'breakout_price': '돌파 가격',
                'cup_depth_pct': '컵 깊이(%)',
                'handle_depth_pct': '핸들 깊이(%)',
                'cup_width_days': '컵 기간(일)',
                'handle_width_days': '핸들 기간(일)',
                'volume_increase_ratio': '거래량 비율',
                'peak_a_date': 'Peak A 날짜',
                'peak_a_price': 'Peak A 가격',
                'valley_b_date': 'Valley B 날짜',
                'valley_b_price': 'Valley B 가격',
                'peak_c_date': 'Peak C 날짜',
                'peak_c_price': 'Peak C 가격',
                'valley_d_date': 'Valley D 날짜',
                'valley_d_price': 'Valley D 가격'
            })
            
            # 수치 데이터 반올림 및 라운딩 가공
            def format_excel_data(row):
                ticker = row['티커']
                is_kr = ticker.endswith('.KS') or ticker.endswith('.KQ')
                
                # 가격 컬럼 포맷팅
                price_cols = ['돌파 가격', 'Peak A 가격', 'Valley B 가격', 'Peak C 가격', 'Valley D 가격']
                for col in price_cols:
                    if col in row and not pd.isna(row[col]):
                        if is_kr:
                            row[col] = int(round(row[col]))
                        else:
                            row[col] = round(row[col], 2)
                
                # 비율 및 수치
                for col in ['컵 깊이(%)', '핸들 깊이(%)', '거래량 비율']:
                    if col in row and not pd.isna(row[col]):
                        row[col] = round(row[col], 2)
                return row
                
            df_excel = df_excel.apply(format_excel_data, axis=1)
            
            # 불필요한 인덱스성 컬럼 제거
            exclude_cols = ['peak_a_idx', 'valley_b_idx', 'peak_c_idx', 'valley_d_idx', 'breakout_idx', 'chart_path']
            df_excel = df_excel.drop(columns=[c for c in exclude_cols if c in df_excel.columns], errors='ignore')
            
            # 엑셀 쓰기
            df_excel.to_excel(writer, index=False, sheet_name='Cup With Handle')
            worksheet = writer.sheets['Cup With Handle']
            
            max_col = worksheet.max_column
            max_row = worksheet.max_row
            if max_row > 0:
                worksheet.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
            
            # 열 너비 자동 맞춤 (기존 대비 10% 축소)
            for col in worksheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val = str(cell.value or '')
                    length = sum(2 if ord(char) > 128 else 1 for char in val)
                    if length > max_len:
                        max_len = length
                # 기존 대비 10% 줄인 너비 설정
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12) * 0.9
                
            # 셀 스타일 및 정렬
            from openpyxl.styles import Alignment, PatternFill, Font
            
            # 1행 헤더 바탕색 및 글꼴 지정 (Dark Navy 배경, 흰색 굵은 글씨)
            header_fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
            header_font = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
            for col_idx in range(1, max_col + 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                
            center_cols = [1, 2, 4, 6, 8, 10]  # 티커(1), 종목명(2), 돌파 감지일(4), 컵 깊이(6), 컵 기간(8), 거래량 비율(10)
            
            for row_idx in range(2, max_row + 1):
                ticker_val = str(worksheet.cell(row=row_idx, column=1).value or '')
                is_kr = ticker_val.endswith('.KS') or ticker_val.endswith('.KQ')
                
                for c_idx in range(1, max_col + 1):
                    cell = worksheet.cell(row=row_idx, column=c_idx)
                    
                    # 1. 정렬 설정 (티커, 종목명, 돌파 감지일, F, H, J열은 가운데 정렬, K, M, O, Q열 및 기타 데이터는 오른쪽 정렬)
                    if c_idx in center_cols:
                        cell.alignment = Alignment(horizontal='center')
                    else:
                        cell.alignment = Alignment(horizontal='right')
                    
                    # 2. 서식 설정
                    # 소수점 데이터 (컵 깊이, 핸들 깊이, 거래량 비율) -> 0.00 형식
                    if c_idx in [6, 7, 10]:
                        cell.number_format = '0.00'
                    # 가격 데이터 -> 한국 주식은 정수(천단위 구분), 미국 주식은 소수점 둘째자리(0.00)
                    elif c_idx in [5, 12, 14, 16, 18]:
                        if is_kr:
                            cell.number_format = '#,##0'
                        else:
                            cell.number_format = '0.00'
                    # 기간 데이터 -> 정수
                    elif c_idx in [8, 9]:
                        cell.number_format = '#,##0'
                        
        excel_data = output.getvalue()
        
        market_code_map = {
            "코스피 (KOSPI)": "KS",
            "코스닥 (KOSDAQ)": "KQ",
            "전체 시장 (KOSPI + KOSDAQ)": "KS&KQ",
            "미국 S&P 500 (US)": "SP",
            "미국 NASDAQ 100 (US)": "NQ"
        }
        market_code = market_code_map.get(st.session_state.market_type_used, "ALL")
        today_str = datetime.datetime.now(KST).strftime('%Y-%m-%d')
        excel_filename = f"CupWithHandle-{market_code}-{today_str}.xlsx"

        # 타이틀 및 엑셀 다운로드 버튼 (동일 라인 양 끝 배치)
        col_title, col_download = st.columns([8, 2], vertical_alignment="bottom")
        with col_title:
            st.markdown(
                f"<div style='font-size: 1.20rem; font-weight: 700; color: #8AB4F8; margin: 10px 0 6px 0; display: flex; align-items: center; gap: 8px;'>"
                f"<span>📋</span> 스크리닝 결과 (총 {len(st.session_state.screened_df)}개 종목)"
                f"</div>",
                unsafe_allow_html=True
            )
        with col_download:
            st.download_button(
                label="📥 엑셀 파일 다운로드",
                data=excel_data,
                file_name=excel_filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        
        # 테이블 소수점 등 출력 포맷 가공
        df_format = st.session_state.screened_df.copy()
        df_format['돌파 가격'] = df_format.apply(lambda r: fmt_curr(r['돌파 가격'], r['티커']), axis=1)
        df_format['컵 깊이(%)'] = df_format['컵 깊이(%)'].map('{:.2f}%'.format)
        df_format['핸들 깊이(%)'] = df_format['핸들 깊이(%)'].map('{:.2f}%'.format)
        df_format['컵 기간(일)'] = df_format['컵 기간(일)'].map('{:,.0f}일'.format)
        df_format['핸들 기간(일)'] = df_format['핸들 기간(일)'].map('{:,.0f}일'.format)
        df_format['거래량 비율'] = df_format['거래량 비율'].map('{:.2f}배'.format)
        
        st.dataframe(df_format, use_container_width=True)
        
        # --- 개별 종목 차트 시각화 영역 ---
        st.markdown("---")
        st.markdown(
            "<div style='font-size: 1.20rem; font-weight: 700; color: #8AB4F8; margin: 20px 0 10px 0; display: flex; align-items: center; gap: 8px;'>"
            "<span>📈</span> 종목별 컵앤핸들 패턴 분석 차트"
            "</div>",
            unsafe_allow_html=True
        )
        
        selected_stock_name = st.selectbox(
            "시각화할 종목을 선택하세요",
            options=st.session_state.screened_df['종목명'].tolist()
        )
        
        if selected_stock_name:
            # 선택된 종목의 정보 추출
            row = st.session_state.raw_screened_df[st.session_state.raw_screened_df['name'] == selected_stock_name].iloc[0]
            ticker = row['symbol']
            
            with st.spinner(f"{selected_stock_name} ({ticker}) 주가 데이터 가져오는 중..."):
                # 차트 작성을 위해 3년치 다운로드
                is_kr = ticker.endswith('.KS') or ticker.endswith('.KQ')
                if is_kr:
                    code = ticker.split('.')[0]
                    start_date = (datetime.datetime.now() - datetime.timedelta(days=1095)).strftime('%Y-%m-%d')
                    df_chart = fdr.DataReader(code, start_date)
                else:
                    df_chart = yf.download(ticker, period="3y", progress=False)
                    if isinstance(df_chart.columns, pd.MultiIndex):
                        df_chart.columns = df_chart.columns.droplevel(1)
                
                df_chart = df_chart.dropna(subset=['Close', 'High', 'Low', 'Volume'])
                if df_chart.index.tz is not None:
                    df_chart.index = df_chart.index.tz_localize(None)
                
            if not df_chart.empty:
                df_chart['SMA_50'] = df_chart['Close'].rolling(window=50).mean()
                df_chart['SMA_150'] = df_chart['Close'].rolling(window=150).mean()
                df_chart['SMA_200'] = df_chart['Close'].rolling(window=200).mean()
                df_chart['Vol_SMA_20'] = df_chart['Volume'].rolling(window=20).mean()
                
                # Plotly 서브플롯 구성 (주가 캔들 + 거래량 바)
                fig = make_subplots(
                    rows=2, cols=1,
                    shared_xaxes=True,
                    vertical_spacing=0.08,
                    row_heights=[0.7, 0.3]
                )
                
                # 1. 캔들스틱 추가
                fig.add_trace(
                    go.Candlestick(
                        x=df_chart.index,
                        open=df_chart['Open'],
                        high=df_chart['High'],
                        low=df_chart['Low'],
                        close=df_chart['Close'],
                        name="주가",
                        increasing_line_color='#EA4335',
                        decreasing_line_color='#4285F4'
                    ),
                    row=1, col=1
                )
                
                # 2. 이동평균선 추가
                fig.add_trace(
                    go.Scatter(x=df_chart.index, y=df_chart['SMA_50'], line=dict(color='#FBBC05', width=1.5), name="50일 MA"),
                    row=1, col=1
                )
                fig.add_trace(
                    go.Scatter(x=df_chart.index, y=df_chart['SMA_150'], line=dict(color='#34A853', width=1.5), name="150일 MA"),
                    row=1, col=1
                )
                fig.add_trace(
                    go.Scatter(x=df_chart.index, y=df_chart['SMA_200'], line=dict(color='#EA4335', width=2), name="200일 MA"),
                    row=1, col=1
                )
                
                # 3. 컵앤핸들 주요 지점 및 저항선 마킹
                # 저항선: Peak A 날짜부터 돌파 날짜까지 Peak C 가격으로 수평선 표시
                try:
                    dt_a = pd.to_datetime(row['peak_a_date'])
                    dt_break = pd.to_datetime(row['breakout_date'])
                    val_c = row['peak_c_price']
                    
                    # 테마에 따라 저항선 색상 결정 (다크 모드: 흰색, 라이트 모드: 검은색)
                    theme_type = "dark"
                    try:
                        if hasattr(st, "context") and st.context.theme:
                            theme_type = getattr(st.context.theme, "type", "dark")
                            if not theme_type:
                                theme_type = st.context.theme.get("type", "dark")
                    except Exception:
                        pass
                    
                    line_color = "#FFFFFF" if theme_type == "dark" else "#000000"
                    
                    fig.add_trace(
                        go.Scatter(
                            x=[dt_a, dt_break],
                            y=[val_c, val_c],
                            mode='lines',
                            line=dict(color=line_color, width=2, dash='dash'),
                            name="컵 저항선 (Peak C)"
                        ),
                        row=1, col=1
                    )
                except Exception:
                    pass
                
                # 극점 마커 그리기 함수
                def add_marker(date_str, price, label, color, symbol='circle'):
                    try:
                        dt = pd.to_datetime(date_str)
                        fig.add_trace(
                            go.Scatter(
                                x=[dt],
                                y=[price],
                                mode='markers',
                                marker=dict(symbol=symbol, size=12, color=color, line=dict(width=1.5, color='black')),
                                name=label,
                                hovertext=f"{label}: {fmt_curr(price, ticker)} ({date_str})",
                                hoverinfo="text"
                            ),
                            row=1, col=1
                        )
                    except Exception:
                        pass
                
                add_marker(row['peak_a_date'], row['peak_a_price'], "Peak A (왼쪽 고점)", "blue", "triangle-down")
                add_marker(row['valley_b_date'], row['valley_b_price'], "Valley B (컵 바닥)", "green", "triangle-up")
                add_marker(row['peak_c_date'], row['peak_c_price'], "Peak C (오른쪽 고점)", "cyan", "triangle-down")
                add_marker(row['valley_d_date'], row['valley_d_price'], "Valley D (핸들 바닥)", "magenta", "triangle-up")
                add_marker(row['breakout_date'], row['breakout_price'], "★ 돌파 (Breakout)", "red", "star")
                
                # 돌파 지점 화살표 어노테이션
                try:
                    dt_break = pd.to_datetime(row['breakout_date'])
                    fig.add_annotation(
                        x=dt_break,
                        y=row['breakout_price'],
                        text="★ 돌파 (Breakout)",
                        showarrow=True,
                        arrowhead=2,
                        arrowsize=1,
                        arrowwidth=2,
                        arrowcolor="#EA4335",
                        ax=0,
                        ay=-40,
                        font=dict(color="#EA4335", size=12, family="Malgun Gothic"),
                        row=1, col=1
                    )
                except Exception:
                    pass
                
                # 4. 거래량 바 추가
                colors = ['#EA4335' if df_chart['Close'].iloc[i] >= df_chart['Open'].iloc[i] else '#4285F4' for i in range(len(df_chart))]
                fig.add_trace(
                    go.Bar(
                        x=df_chart.index,
                        y=df_chart['Volume'],
                        marker_color=colors,
                        name="거래량"
                    ),
                    row=2, col=1
                )
                
                # 거래량 20MA
                fig.add_trace(
                    go.Scatter(
                        x=df_chart.index,
                        y=df_chart['Vol_SMA_20'],
                        line=dict(color='#5F6368', width=1.5),
                        name="20일 거래량 MA"
                    ),
                    row=2, col=1
                )
                
                # 화폐 단위 동적 결정
                is_us_stock = not (ticker.endswith('.KS') or ticker.endswith('.KQ'))
                currency_symbol = '$' if is_us_stock else '원'
                tick_format = ',.2f' if is_us_stock else ',.0f'
                
                fig.update_layout(
                    template="plotly_dark",
                    paper_bgcolor=STANDARD_CHART_THEME['paper_bgcolor'],
                    plot_bgcolor=STANDARD_CHART_THEME['plot_bgcolor'],
                    title=dict(
                        text=f"<b>📈 {selected_stock_name} ({ticker}) 'Cup with Handle' 분석 차트</b>",
                        font=dict(color="#F8FAFC", size=16)
                    ),
                    yaxis_title=f"주가 ({currency_symbol})",
                    yaxis2_title="거래량 (주)",
                    xaxis_rangeslider_visible=False,
                    height=700,
                    margin=dict(l=50, r=50, t=80, b=50),
                    showlegend=True,
                    legend=dict(
                        orientation="h",
                        y=1.08,
                        xanchor="right",
                        x=1,
                        bgcolor="rgba(30, 41, 59, 0.85)",
                        bordercolor="#334155",
                        borderwidth=1,
                        font=dict(color="#F8FAFC", size=11)
                    ),
                    hovermode="x unified"
                )
                
                fig.update_xaxes(
                    tickformat="%Y-%m-%d",
                    hoverformat="%Y-%m-%d",
                    gridcolor="#334155", 
                    linecolor="#475569", 
                    tickfont=dict(color="#cbd5e1")
                )
                fig.update_yaxes(tickformat=tick_format, row=1, col=1, gridcolor="#334155", linecolor="#475569", tickfont=dict(color="#cbd5e1"))
                fig.update_yaxes(tickformat=",.0f", row=2, col=1, gridcolor="#334155", linecolor="#475569", tickfont=dict(color="#cbd5e1"))
                
                # 주말 및 공휴일 공백 제거 (5일 주기 끊김 및 0값 방지)
                dt_all = pd.date_range(start=df_chart.index[0], end=df_chart.index[-1], freq='B')
                existing_dates = set(pd.to_datetime(df_chart.index).normalize())
                holidays = [d.strftime('%Y-%m-%d') for d in dt_all if d.normalize() not in existing_dates]
                rbreaks = [dict(bounds=["sat", "mon"])]
                if holidays:
                    rbreaks.append(dict(values=holidays))
                fig.update_xaxes(rangebreaks=rbreaks)
                
                st.plotly_chart(fig, use_container_width=True)
                
                # 상세 분석 정보 카드 출력
                st.markdown(f"""
                <div class="metric-card">
                    <div style="font-size: 1.00rem; font-weight: 600; color: #E2E8F0; margin-bottom: 10px; display: flex; align-items: center; gap: 6px;">
                        <span>💡</span> {selected_stock_name} ({ticker}) 컵앤핸들 상세 패턴 통계
                    </div>
                    <ul>
                        <li><b>돌파 가격:</b> {fmt_curr(row['breakout_price'], ticker)} (돌파 감지일: {row['breakout_date']})</li>
                        <li><b>돌파 시점 거래량 비율:</b> <span style="color:#EA4335; font-weight:bold;">{row['volume_increase_ratio']:.2f}배</span> (20일 평균 거래량 대비)</li>
                        <li><b>컵 깊이(조정폭):</b> {row['cup_depth_pct']:.1f}% (기간: {row['cup_width_days']} 거래일)</li>
                        <li><b>핸들 깊이(조정폭):</b> {row['handle_depth_pct']:.1f}% (기간: {row['handle_width_days']} 거래일)</li>
                        <li><b>세부 극점 가격 정보:</b>
                            <ul>
                                <li>왼쪽 고점(Peak A): {fmt_curr(row['peak_a_price'], ticker)} ({row['peak_a_date']})</li>
                                <li>컵 최저점(Valley B): {fmt_curr(row['valley_b_price'], ticker)} ({row['valley_b_date']})</li>
                                <li>오른쪽 고점(Peak C): {fmt_curr(row['peak_c_price'], ticker)} ({row['peak_c_date']})</li>
                                <li>핸들 최저점(Valley D): {fmt_curr(row['valley_d_price'], ticker)} ({row['valley_d_date']})</li>
                            </ul>
                        </li>
                    </ul>
                </div>
                """, unsafe_allow_html=True)
                
else:
    st.info("👈 왼쪽 사이드바에서 대상 시장 및 파라미터를 설정한 후 '스크리닝 시작' 버튼을 눌러주세요.")

st.markdown("---")
st.markdown("<div style='text-align: center; color: #64748b; font-size: 0.8rem; margin-top: 8px; margin-bottom: 24px; line-height: 1.6;'>⚠️ 본 서비스에서 제공하는 모든 정보는 투자 참고용이며, 투자의 최종 결정과 책임은 투자자 본인에게 있습니다.</div>", unsafe_allow_html=True)
